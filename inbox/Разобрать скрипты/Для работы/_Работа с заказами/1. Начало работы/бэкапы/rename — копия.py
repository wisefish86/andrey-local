import logging
import os
import pathlib
import re
import shutil
import subprocess
import sys
import zipfile

import openpyxl
import pyperclip

# ================= НАСТРОЙКИ =================
CLIENT_CELL = 'D4'                # ячейка с заказчиком
MANAGER_CELL = 'D5'               # ячейка с менеджером
EMAIL_FILENAME = "- Emails для согласования.txt"
ITEMS_FILENAME = "- Позиции заказа.txt"
LOG_FILENAME = "rename.log"       # файл лога (в папке скрипта)
FILES_FOLDER_NAME = "файлы"

ORDERS_DIR = pathlib.Path(r"D:\YandexDisk\Проекты\Работа в Компас\Заказы")
TEMPLATE_FOLDER = ORDERS_DIR / "● [копировать] Заказ 00000 (ЗАКАЗЧИК) - МЕНЕДЖЕР"

# ---------- Настройки для позиций заказа ----------
ITEMS_START_ROW = 8               # строка первой позиции
ITEM_NUM_COL = 'C'                # колонка с номером
ITEM_NAME_COL = 'D'               # колонка с наименованием
ITEM_QTY_COL = 'E'                # колонка с количеством
ITEM_COMMENT_COL = 'F'            # колонка с комментарием
# ====================================================


def clean_text(raw):
    """
    Удаляет подстроки вида _xNNNN_ (например, _x000D_), но оставляет реальные переносы строк.
    """
    if raw is None:
        return ""
    text = str(raw)
    text = re.sub(r'_x[0-9A-Fa-f]{4}_', '', text)
    return text.strip()


script_path = pathlib.Path(__file__).resolve()
script_dir = script_path.parent
log_file = script_dir / LOG_FILENAME

logger = logging.getLogger('rename')
logger.setLevel(logging.DEBUG)
logger.handlers.clear()

# Лог каждый запуск создается заново, без накопления старого мусора.
file_handler = logging.FileHandler(log_file, encoding='utf-8', mode='w')
file_handler.setLevel(logging.DEBUG)
formatter = logging.Formatter('%(asctime)s - %(levelname)s - %(message)s', datefmt='%Y-%m-%d %H:%M:%S')
file_handler.setFormatter(formatter)
logger.addHandler(file_handler)

console_handler = logging.StreamHandler()
console_handler.setLevel(logging.INFO)
console_handler.setFormatter(formatter)
logger.addHandler(console_handler)

logger.info("=== Python folder rename started ===")


def fail(message):
    logger.error(message)
    input("Press Enter...")
    sys.exit(1)


def normalize_arg(raw_arg):
    folder_str = raw_arg.strip().strip('"').strip()
    path = pathlib.Path(folder_str)
    if not path.is_absolute():
        path = (script_dir / path).resolve()
    return path


def unique_path(path):
    if not path.exists():
        return path

    parent = path.parent

    if path.is_dir() or not path.suffix:
        stem = path.name
        suffix = ""
    else:
        stem = path.stem
        suffix = path.suffix

    for index in range(1, 1000):
        candidate = parent / f"{stem} ({index}){suffix}"
        if not candidate.exists():
            return candidate

    raise FileExistsError(f"Cannot create unique name for {path}")


def is_archive(path):
    return path.is_file() and path.suffix.lower() in {
        ".zip", ".7z", ".rar", ".tar", ".gz", ".tgz", ".bz2", ".tbz2", ".xz", ".txz"
    }


def extract_archive(archive_path, destination):
    suffix = archive_path.suffix.lower()

    if suffix == ".zip":
        with zipfile.ZipFile(archive_path) as archive:
            archive.extractall(destination)
        return

    try:
        shutil.unpack_archive(str(archive_path), str(destination))
        return
    except (shutil.ReadError, ValueError):
        pass

    seven_zip = shutil.which("7z") or shutil.which("7za") or shutil.which("7zr")
    if seven_zip:
        subprocess.run(
            [seven_zip, "x", str(archive_path), f"-o{destination}", "-y"],
            check=True,
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
        )
        return

    raise RuntimeError(
        f"Cannot unpack {archive_path.name}. Для .rar/.7z нужен установленный 7-Zip в PATH."
    )


def move_file_or_folder(source, destination_dir):
    destination = unique_path(destination_dir / source.name)
    shutil.move(str(source), str(destination))
    logger.info(f"Moved: {source.name} -> {destination}")


def is_order_excel(path):
    return path.is_file() and path.name.startswith('Заказ') and path.suffix.lower() in ['.xlsx', '.xlsm']


def should_stay_in_order_root(item):
    if item.is_dir():
        return False
    if is_order_excel(item):
        return True
    if item.suffix.lower() == ".cdr":
        return True
    if item.suffix.lower() in [".txt", ".url"] and item.name.startswith("-"):
        return True
    return False


def move_customer_files_to_files_folder(order_folder):
    files_folder = order_folder / FILES_FOLDER_NAME
    files_folder.mkdir(exist_ok=True)

    for item in order_folder.iterdir():
        if item == files_folder:
            continue
        if should_stay_in_order_root(item):
            continue
        move_file_or_folder(item, files_folder)

    return files_folder


def find_cdr_file(order_folder):
    search_dirs = [order_folder, order_folder / FILES_FOLDER_NAME]
    for search_dir in search_dirs:
        if not search_dir.exists() or not search_dir.is_dir():
            continue
        for f in search_dir.iterdir():
            if f.is_file() and f.suffix.lower() == ".cdr":
                return f
    return None


def create_order_folder_from_inputs(input_paths):
    if not TEMPLATE_FOLDER.exists() or not TEMPLATE_FOLDER.is_dir():
        fail(f"Template folder does not exist: {TEMPLATE_FOLDER}")

    ORDERS_DIR.mkdir(parents=True, exist_ok=True)
    new_folder = unique_path(ORDERS_DIR / TEMPLATE_FOLDER.name)
    shutil.copytree(TEMPLATE_FOLDER, new_folder)
    logger.info(f"Template copied to: {new_folder}")

    for input_path in input_paths:
        if not input_path.exists():
            fail(f"Dragged path does not exist: {input_path}")

        if is_archive(input_path):
            logger.info(f"Extracting archive: {input_path.name}")
            extract_archive(input_path, new_folder)
            try:
                input_path.unlink()
                logger.info(f"Archive deleted: {input_path}")
            except Exception as e:
                logger.error(f"Cannot delete archive {input_path}: {e}")
        else:
            move_file_or_folder(input_path, new_folder)

    return new_folder


def get_working_folder(args):
    if not args:
        fail("No target folder, archive or files passed. Use bat file.")

    paths = [normalize_arg(arg) for arg in args]
    logger.debug(f"Raw argv: {[repr(arg) for arg in args]}")
    logger.info(f"Received paths: {len(paths)}")

    if len(paths) == 1 and paths[0].is_dir():
        return paths[0]

    return create_order_folder_from_inputs(paths)


folder = get_working_folder(sys.argv[1:])

logger.info(f"Working folder: {folder}")

if not folder.exists() or not folder.is_dir():
    fail("Target folder does not exist or is not a directory.")

files_folder = move_customer_files_to_files_folder(folder)
logger.info(f"Files folder: {files_folder}")

# --- Поиск Excel-файла, начинающегося с "Заказ" ---
excel_file = None
for f in folder.iterdir():
    if is_order_excel(f):
        excel_file = f
        break

if not excel_file:
    fail("No Excel file starting with 'Заказ' found.")

logger.info(f"Excel file: {excel_file.name}")

wb = openpyxl.load_workbook(excel_file, data_only=True)
ws = wb.active

# Чтение данных из ячеек с очисткой
line = clean_text(ws['B2'].value)
client = clean_text(ws[CLIENT_CELL].value)
manager_raw = clean_text(ws[MANAGER_CELL].value)

# Очистка менеджера от префикса "Менеджер:"
if manager_raw:
    if ':' in manager_raw:
        manager = manager_raw.split(':', 1)[-1].strip()
    else:
        manager = manager_raw
else:
    manager = ""

logger.info(f"B2: {line}")
logger.info(f"CLIENT ({CLIENT_CELL}): {client}")
logger.info(f"MANAGER ({MANAGER_CELL}): {manager}")

# Извлечение номера и даты из B2
m = re.search(r"\D*(\d+)\D+(\d{2}\.\d{2}\.\d{4})", line)
if not m:
    fail("ERROR parsing B2.")

order_number = m.group(1)
order_date = m.group(2)

# Формирование нового имени папки
base_name = f"Заказ {order_number} от {order_date} ({client})"
if manager:
    new_name_folder = f"{base_name} - {manager}"
else:
    new_name_folder = base_name

# Удаление недопустимых символов для Windows
invalid = '<>:"/\\|?*'
for ch in invalid:
    new_name_folder = new_name_folder.replace(ch, "_")
new_name_folder = new_name_folder.strip()

logger.info(f"New folder name: {new_name_folder}")

# Формирование текста для буфера обмена
clipboard_lines = [
    f"Заказ {order_number} от {order_date}",
    f"Заказчик: {client}",
    f"Менеджер: {manager}"
]
clipboard_text = "\r\n".join(clipboard_lines)

# Копирование в буфер
try:
    pyperclip.copy(clipboard_text)
    logger.info("Copied to clipboard:\n" + clipboard_text.replace('\r', ''))
except Exception as e:
    logger.error(f"Clipboard error: {e}")

# --- Переименование папки ---
parent = folder.parent
target_folder = parent / new_name_folder

if target_folder == folder:
    logger.info("Folder already correctly named.")
    work_folder = folder
elif target_folder.exists():
    target_folder = unique_path(target_folder)
    try:
        folder.rename(target_folder)
        work_folder = target_folder
        logger.info(f"Folder renamed with unique suffix: {target_folder.name}")
    except PermissionError:
        logger.error("Folder busy, cannot rename. Working inside current folder.")
        work_folder = folder
else:
    try:
        folder.rename(target_folder)
        work_folder = target_folder
        logger.info("Folder renamed.")
    except PermissionError:
        logger.error("Folder busy, cannot rename. Working inside current folder.")
        work_folder = folder

logger.info(f"Work folder now: {work_folder}")
files_folder = work_folder / FILES_FOLDER_NAME
files_folder.mkdir(exist_ok=True)
final_order_name = work_folder.name

# --- Открытие папки в Проводнике ---
try:
    os.startfile(str(work_folder))
    logger.info("Folder opened in Explorer.")
except Exception as e:
    logger.error(f"Cannot open folder: {e}")

# --- Извлечение email'ов и запись в файл ---
email_pattern = r'[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}'
found_emails = []

for row in range(1, ws.max_row + 1):
    cell = ws[f'B{row}']
    if cell.value and "Контакты для согласования макетов" in str(cell.value):
        cell_text = clean_text(str(cell.value))
        found_emails = re.findall(email_pattern, cell_text)
        if found_emails:
            logger.info(f"Found emails: {', '.join(found_emails)}")
        break

email_file_path = work_folder / EMAIL_FILENAME
if found_emails:
    try:
        with open(email_file_path, 'w', encoding='utf-8') as f:
            f.write(', '.join(found_emails))
        logger.info(f"Emails saved to: {EMAIL_FILENAME}")
    except Exception as e:
        logger.error(f"Error writing email file: {e}")
else:
    logger.info("No emails found.")


# --- Извлечение позиций заказа (жестко по настройкам) ---
def extract_items_fixed(worksheet):
    """
    Собирает позиции, начиная с ITEMS_START_ROW, пока в колонке ITEM_NAME_COL есть значение.
    Формат:
        основная строка: номер - наименование - количество шт.
        если есть комментарий: новая строка (в круглых скобках, с сохранением внутренних переносов).
    Между позициями позже будут добавлены пустые строки.
    """
    items = []
    for row in range(ITEMS_START_ROW, worksheet.max_row + 1):
        name_cell = worksheet[f'{ITEM_NAME_COL}{row}'].value
        if name_cell is None:
            break
        name = clean_text(name_cell)
        if name == '':
            break

        num_cell = worksheet[f'{ITEM_NUM_COL}{row}'].value
        item_num = clean_text(num_cell)

        qty_cell = worksheet[f'{ITEM_QTY_COL}{row}'].value
        quantity = clean_text(qty_cell)

        comment_cell = worksheet[f'{ITEM_COMMENT_COL}{row}'].value
        comment = clean_text(comment_cell)

        # Основная строка
        line = f"{item_num} - {name} - {quantity} шт."
        if comment:
            # Комментарий с новой строки, в круглых скобках
            line += f"\n({comment})"
        items.append(line)
        logger.debug(f"Добавлена позиция: {line}")

    logger.info(f"Собрано позиций: {len(items)}")
    return items


items = extract_items_fixed(ws)
items_file_path = work_folder / ITEMS_FILENAME
if items:
    try:
        # Формируем заголовок из основных данных (как в буфере обмена)
        header_lines = [
            f"Заказ {order_number} от {order_date}",
            f"Заказчик: {client}",
            f"Менеджер: {manager}"
        ]
        header = "\n".join(header_lines)
        content = header + "\n\n\n" + "\n\n".join(items)
        with open(items_file_path, 'w', encoding='utf-8') as f:
            f.write(content)
        logger.info(f"Items saved to: {ITEMS_FILENAME} (найдено {len(items)} позиций)")
    except Exception as e:
        logger.error(f"Error writing items file: {e}")
else:
    logger.warning("No items found in the table.")
    if items_file_path.exists():
        try:
            items_file_path.unlink()
            logger.info(f"Removed old {ITEMS_FILENAME} (no items).")
        except Exception as e:
            logger.error(f"Error removing old items file: {e}")

# Актуализация пути к Excel
excel_path = work_folder / excel_file.name

# --- Переименование CDR-файла (с заменой) ---
cdr_file = None
cdr_file = find_cdr_file(work_folder)

if not cdr_file:
    logger.warning("No .cdr found.")
    try:
        os.startfile(str(excel_path))
        logger.info("Excel opened.")
    except Exception as e:
        logger.error(f"Cannot open Excel: {e}")
else:
    new_cdr = cdr_file.parent / f"{final_order_name}.cdr"
    if cdr_file == new_cdr:
        logger.info("CDR already has correct name.")
        target_cdr = cdr_file
    else:
        try:
            os.replace(str(cdr_file), str(new_cdr))
            logger.info("CDR renamed (replaced if existed).")
            target_cdr = new_cdr
        except PermissionError:
            logger.error("CDR file busy, opening original.")
            target_cdr = cdr_file
        except OSError as e:
            logger.error(f"Error renaming CDR: {e}")
            target_cdr = cdr_file

    # --- Открытие CorelDRAW ---
    try:
        os.startfile(str(target_cdr))
        logger.info("CorelDRAW opened.")
    except Exception as e:
        logger.error(f"Cannot open CDR: {e}")

# --- Открытие Excel (всегда, даже если нет CDR) ---
try:
    os.startfile(str(excel_path))
    logger.info("Excel opened.")
except Exception as e:
    logger.error(f"Cannot open Excel: {e}")

# --- Открытие файла с позициями заказа (если создан) ---
if items_file_path.exists():
    try:
        os.startfile(str(items_file_path))
        logger.info(f"Opened {ITEMS_FILENAME}")
    except Exception as e:
        logger.error(f"Cannot open items file: {e}")

logger.info("DONE.")
