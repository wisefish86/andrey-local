import logging
import os
import pathlib
import re
import shutil
import subprocess
import sys
import zipfile

import openpyxl
import pyperclip

# ================= НАСТРОЙКИ =================
CLIENT_CELL = 'D4'                # ячейка с заказчиком
MANAGER_CELL = 'D5'               # ячейка с менеджером
EMAIL_FILENAME = "- Emails для согласования.txt"
ITEMS_FILENAME = "- Позиции заказа.txt"
LOG_FILENAME = "rename.log"       # файл лога (в папке скрипта)
FILES_FOLDER_NAME = "файлы"
TXT_MANAGER = "Наталия Миронова"
TXT_SOURCE_FILENAME = "заказ от Мироновой.txt"
SCRIPT_VERSION = "2026-07-24-txt-manager-v4"

ORDERS_DIR = pathlib.Path(r"D:\YandexDisk\Проекты\Работа в Компас\Заказы")
TEMPLATE_FOLDER = ORDERS_DIR / "● [копировать] Заказ 00000 (ЗАКАЗЧИК) - МЕНЕДЖЕР"

# ---------- Настройки для позиций заказа ----------
ITEMS_START_ROW = 8               # строка первой позиции
ITEM_NUM_COL = 'C'                # колонка с номером
ITEM_NAME_COL = 'D'               # колонка с наименованием
ITEM_QTY_COL = 'E'                # колонка с количеством
ITEM_COMMENT_COL = 'F'            # колонка с комментарием
# ====================================================

EMAIL_PATTERN = r'[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}'


def clean_text(raw):
    """
    Удаляет подстроки вида _xNNNN_ (например, _x000D_),
    но оставляет реальные переносы строк.
    """
    if raw is None:
        return ""
    text = str(raw)
    text = re.sub(r'_x[0-9A-Fa-f]{4}_', '', text)
    return text.strip()


script_path = pathlib.Path(__file__).resolve()
script_dir = script_path.parent
log_file = script_dir / LOG_FILENAME

logger = logging.getLogger('rename')
logger.setLevel(logging.DEBUG)
logger.handlers.clear()

# Лог каждый запуск создается заново, без накопления старого мусора.
file_handler = logging.FileHandler(log_file, encoding='utf-8', mode='w')
file_handler.setLevel(logging.DEBUG)
formatter = logging.Formatter(
    '%(asctime)s - %(levelname)s - %(message)s',
    datefmt='%Y-%m-%d %H:%M:%S'
)
file_handler.setFormatter(formatter)
logger.addHandler(file_handler)

console_handler = logging.StreamHandler()
console_handler.setLevel(logging.INFO)
console_handler.setFormatter(formatter)
logger.addHandler(console_handler)

logger.info("=== Python folder rename started ===")
logger.info(f"SCRIPT VERSION: {SCRIPT_VERSION}")
logger.info(f"SCRIPT PATH: {script_path}")


def fail(message):
    logger.error(message)
    input("Press Enter...")
    sys.exit(1)


def normalize_arg(raw_arg):
    folder_str = raw_arg.strip().strip('"').strip()
    path = pathlib.Path(folder_str)
    if not path.is_absolute():
        path = (script_dir / path).resolve()
    return path


def unique_path(path):
    if not path.exists():
        return path

    parent = path.parent

    if path.is_dir() or not path.suffix:
        stem = path.name
        suffix = ""
    else:
        stem = path.stem
        suffix = path.suffix

    for index in range(1, 1000):
        candidate = parent / f"{stem} ({index}){suffix}"
        if not candidate.exists():
            return candidate

    raise FileExistsError(f"Cannot create unique name for {path}")


def is_archive(path):
    return path.is_file() and path.suffix.lower() in {
        ".zip", ".7z", ".rar", ".tar", ".gz", ".tgz", ".bz2", ".tbz2", ".xz", ".txz"
    }


def extract_archive(archive_path, destination):
    suffix = archive_path.suffix.lower()

    if suffix == ".zip":
        with zipfile.ZipFile(archive_path) as archive:
            archive.extractall(destination)
        return

    try:
        shutil.unpack_archive(str(archive_path), str(destination))
        return
    except (shutil.ReadError, ValueError):
        pass

    seven_zip = shutil.which("7z") or shutil.which("7za") or shutil.which("7zr")
    if seven_zip:
        subprocess.run(
            [seven_zip, "x", str(archive_path), f"-o{destination}", "-y"],
            check=True,
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
        )
        return

    raise RuntimeError(
        f"Cannot unpack {archive_path.name}. Для .rar/.7z нужен установленный 7-Zip в PATH."
    )


def move_file_or_folder(source, destination_dir):
    destination = unique_path(destination_dir / source.name)
    shutil.move(str(source), str(destination))
    logger.info(f"Moved: {source.name} -> {destination}")


def is_order_excel(path):
    return (
        path.is_file()
        and path.name.startswith('Заказ')
        and path.suffix.lower() in ['.xlsx', '.xlsm']
    )


def should_stay_in_order_root(item):
    if item.is_dir():
        return False
    if is_order_excel(item):
        return True
    if item.suffix.lower() == ".cdr":
        return True
    if item.suffix.lower() in [".txt", ".url"] and item.name.startswith("-"):
        return True
    return False


def move_customer_files_to_files_folder(order_folder):
    files_folder = order_folder / FILES_FOLDER_NAME
    files_folder.mkdir(exist_ok=True)

    for item in order_folder.iterdir():
        if item == files_folder:
            continue
        if should_stay_in_order_root(item):
            continue
        move_file_or_folder(item, files_folder)

    return files_folder


def rename_source_txt(source_txt_file):
    """Переименовывает входной TXT от Мироновой, сохраняя его расширение и папку."""
    target = source_txt_file.with_name(TXT_SOURCE_FILENAME)

    if source_txt_file == target:
        logger.info(f"Source TXT already has correct name: {TXT_SOURCE_FILENAME}")
        return source_txt_file

    try:
        # Не затираем другой файл с таким именем: сначала сохраняем его под уникальным именем.
        if target.exists():
            previous = unique_path(target.with_name(
                f"{target.stem} - предыдущий{target.suffix}"
            ))
            target.rename(previous)
            logger.warning(f"Existing {TXT_SOURCE_FILENAME} moved to: {previous.name}")

        source_txt_file.rename(target)
        logger.info(f"Source TXT renamed: {source_txt_file.name} -> {TXT_SOURCE_FILENAME}")
        return target
    except Exception as e:
        logger.error(f"Cannot rename source TXT {source_txt_file.name}: {e}")
        return source_txt_file


def ensure_source_txt_in_files_folder(source_txt_file, files_folder):
    """Гарантирует, что исходный TXT заказа находится внутри папки «файлы»."""
    try:
        source_txt_file.relative_to(files_folder)
        logger.info(f"Source TXT already inside files folder: {source_txt_file}")
        return source_txt_file
    except ValueError:
        pass

    destination = unique_path(files_folder / source_txt_file.name)
    shutil.move(str(source_txt_file), str(destination))
    logger.info(f"Source TXT moved to files folder: {destination}")
    return destination


def has_additional_files(files_folder, source_txt_file=None):
    """
    Проверяет, есть ли в папке «файлы» что-либо помимо исходного TXT заказа.
    Учитываются файлы и вложенные папки. Это сделано намеренно жестко:
    если есть любой дополнительный материал, открывается папка «файлы».
    """
    source_resolved = None
    if source_txt_file is not None:
        try:
            source_resolved = source_txt_file.resolve()
        except Exception:
            source_resolved = source_txt_file

    contents = []
    for path in files_folder.iterdir():
        try:
            current = path.resolve()
        except Exception:
            current = path

        # Единственный входной TXT заказа не считается дополнительным вложением.
        if source_resolved is not None and current == source_resolved:
            logger.debug(f"Source TXT excluded from additional content: {path}")
            continue

        contents.append(path)

    if contents:
        logger.info(
            "Additional content found in files folder: "
            + "; ".join(str(path) for path in contents)
        )
        return True

    logger.info("No additional content found in files folder besides source TXT.")
    return False


def find_cdr_file(order_folder):
    search_dirs = [order_folder, order_folder / FILES_FOLDER_NAME]
    for search_dir in search_dirs:
        if not search_dir.exists() or not search_dir.is_dir():
            continue
        for f in search_dir.iterdir():
            if f.is_file() and f.suffix.lower() == ".cdr":
                return f
    return None


def create_order_folder_from_inputs(input_paths):
    if not TEMPLATE_FOLDER.exists() or not TEMPLATE_FOLDER.is_dir():
        fail(f"Template folder does not exist: {TEMPLATE_FOLDER}")

    ORDERS_DIR.mkdir(parents=True, exist_ok=True)
    new_folder = unique_path(ORDERS_DIR / TEMPLATE_FOLDER.name)
    shutil.copytree(TEMPLATE_FOLDER, new_folder)
    logger.info(f"Template copied to: {new_folder}")

    for input_path in input_paths:
        if not input_path.exists():
            fail(f"Dragged path does not exist: {input_path}")

        if is_archive(input_path):
            logger.info(f"Extracting archive: {input_path.name}")
            extract_archive(input_path, new_folder)
            try:
                input_path.unlink()
                logger.info(f"Archive deleted: {input_path}")
            except Exception as e:
                logger.error(f"Cannot delete archive {input_path}: {e}")
        else:
            move_file_or_folder(input_path, new_folder)

    return new_folder


def get_working_folder(args):
    if not args:
        fail("No target folder, archive or files passed. Use bat file.")

    paths = [normalize_arg(arg) for arg in args]
    logger.debug(f"Raw argv: {[repr(arg) for arg in args]}")
    logger.info(f"Received paths: {len(paths)}")

    if len(paths) == 1 and paths[0].is_dir():
        return paths[0]

    return create_order_folder_from_inputs(paths)


# ================= ЧТЕНИЕ TXT =================

def read_txt_with_encoding(path):
    """
    Читает TXT в распространенных кодировках.
    Возвращает кортеж: (текст, название кодировки).
    """
    raw = path.read_bytes()

    if raw.startswith((b'\xff\xfe', b'\xfe\xff')):
        encodings = ['utf-16', 'utf-8-sig', 'utf-8', 'cp1251']
    else:
        encodings = ['utf-8-sig', 'utf-8', 'cp1251']

    last_error = None
    for encoding in encodings:
        try:
            return raw.decode(encoding), encoding
        except UnicodeDecodeError as e:
            last_error = e

    raise UnicodeError(f"Не удалось определить кодировку TXT: {path.name}. {last_error}")


def looks_like_order_txt(text):
    """Проверяет, похож ли TXT на файл с данными заказа."""
    has_order = re.search(r'(?im)^\s*заказ\s*(?:№|N)?\s*:?\s*\d+', text) is not None
    has_client = re.search(r'(?im)^\s*заказчик\s*:', text) is not None
    has_items = re.search(r'(?im)^\s*состав\s+заказа\s*:', text) is not None
    return has_order and has_client and has_items


def find_order_txt(order_folder):
    """
    Ищет TXT с данными заказа по содержимому, а не по имени.
    Поиск идет в корне заказа, в папке «файлы» и во вложенных папках.
    Служебные выходные TXT исключаются.
    """
    ignored_names = {EMAIL_FILENAME.lower(), ITEMS_FILENAME.lower()}
    candidates = []

    for path in order_folder.rglob('*'):
        if not path.is_file() or path.suffix.lower() != '.txt':
            continue
        if path.name.lower() in ignored_names:
            continue

        try:
            text, encoding = read_txt_with_encoding(path)
        except Exception as e:
            logger.warning(f"Cannot read TXT {path}: {e}")
            continue

        if looks_like_order_txt(text):
            candidates.append((path, text, encoding))

    if not candidates:
        return None

    if len(candidates) > 1:
        names = '; '.join(str(item[0]) for item in candidates)
        fail(f"Найдено несколько TXT с данными заказа: {names}")

    return candidates[0]


def remove_inn_from_client(raw_client):
    """
    Удаляет ИНН из строки заказчика. ИНН отдельно не сохраняется.

    Поддерживаются, например:
        (ИНН: 7838490236) Сибгазполимер
        Сибгазполимер (ИНН: 7838490236)
        Сибгазполимер, ИНН 7838490236
    """
    client = clean_text(raw_client)
    client = re.sub(r'\(\s*ИНН\s*:?\s*[^)]*\)', '', client, flags=re.IGNORECASE)
    client = re.sub(r'\bИНН\s*:?\s*\d+\b', '', client, flags=re.IGNORECASE)
    client = re.sub(r'\s{2,}', ' ', client)
    return client.strip(' \t,;:—–-()')


def unique_emails(text):
    """Находит email и удаляет повторы без изменения исходного порядка."""
    result = []
    seen = set()

    for email in re.findall(EMAIL_PATTERN, text):
        normalized = email.lower()
        if normalized in seen:
            continue
        seen.add(normalized)
        result.append(email)

    return result


def parse_txt_items(text):
    """
    Извлекает позиции после заголовка «Состав заказа:».
    Поддерживает позиции, перенесенные на несколько строк.
    """
    section_match = re.search(r'(?im)^\s*состав\s+заказа\s*:\s*$', text)
    if not section_match:
        return []

    section = text[section_match.end():]
    raw_items = []
    current_num = None
    current_parts = []

    def flush_current():
        if current_num is not None and current_parts:
            joined = ' '.join(part.strip() for part in current_parts if part.strip())
            joined = re.sub(r'\s+', ' ', joined).strip()
            raw_items.append((current_num, joined))

    for raw_line in section.splitlines():
        line = clean_text(raw_line)
        item_match = re.match(r'^\s*(\d+)\s*[.)]\s*(.*)$', line)

        if item_match:
            flush_current()
            current_num = item_match.group(1)
            current_parts = [item_match.group(2)]
        elif current_num is not None and line:
            current_parts.append(line)

    flush_current()

    items = []
    for item_num, raw_item in raw_items:
        # Количество берем из конца строки, поэтому дефисы внутри названия не мешают.
        qty_match = re.match(
            r'^(.*?)\s+-\s+(\d+(?:[.,]\d+)?)\s*шт\.?\s*$',
            raw_item,
            flags=re.IGNORECASE
        )

        if qty_match:
            name = clean_text(qty_match.group(1))
            quantity = clean_text(qty_match.group(2))
            line = f"{item_num} - {name} - {quantity} шт."
        else:
            # Не теряем позицию даже при неожиданном формате количества.
            line = f"{item_num} - {raw_item}"
            logger.warning(
                f"Не удалось отдельно распознать количество в позиции {item_num}: {raw_item}"
            )

        items.append(line)
        logger.debug(f"Добавлена TXT-позиция: {line}")

    logger.info(f"Собрано TXT-позиций: {len(items)}")
    return items


def parse_order_txt(text):
    """Разбирает TXT и возвращает данные, нужные остальному скрипту."""
    order_match = re.search(
        r'(?im)^\s*заказ\s*(?:№|N)?\s*:?\s*(\d+)(?:\s+от\s+(\d{2}\.\d{2}\.\d{4}))?',
        text
    )
    if not order_match:
        raise ValueError("не найден номер заказа")

    order_number = order_match.group(1)
    order_date = order_match.group(2) or ""

    if not order_date:
        date_match = re.search(
            r'(?im)^\s*дата(?:\s+заказа)?\s*:\s*(\d{2}\.\d{2}\.\d{4})\s*$',
            text
        )
        if date_match:
            order_date = date_match.group(1)

    client_match = re.search(r'(?im)^\s*заказчик\s*:\s*(.*?)\s*$', text)
    if not client_match:
        raise ValueError("не найден заказчик")

    client = remove_inn_from_client(client_match.group(1))
    if not client:
        raise ValueError("не удалось выделить название заказчика")

    # TXT-заказы передает только Наталия Миронова, поэтому менеджер фиксированный.
    manager = TXT_MANAGER

    emails = unique_emails(text)
    items = parse_txt_items(text)
    if not items:
        raise ValueError("не найдены позиции заказа")

    return {
        'order_number': order_number,
        'order_date': order_date,
        'client': client,
        'manager': manager,
        'emails': emails,
        'items': items,
    }


# ================= ЧТЕНИЕ EXCEL =================

def extract_excel_emails(worksheet):
    found_emails = []

    for row in range(1, worksheet.max_row + 1):
        cell = worksheet[f'B{row}']
        if cell.value and "Контакты для согласования макетов" in str(cell.value):
            cell_text = clean_text(str(cell.value))
            found_emails = unique_emails(cell_text)
            if found_emails:
                logger.info(f"Found emails: {', '.join(found_emails)}")
            break

    return found_emails


def extract_items_fixed(worksheet):
    """
    Собирает позиции, начиная с ITEMS_START_ROW, пока в колонке ITEM_NAME_COL есть значение.
    Формат:
        основная строка: номер - наименование - количество шт.
        если есть комментарий: новая строка (в круглых скобках, с сохранением внутренних переносов).
    Между позициями позже будут добавлены пустые строки.
    """
    items = []
    for row in range(ITEMS_START_ROW, worksheet.max_row + 1):
        name_cell = worksheet[f'{ITEM_NAME_COL}{row}'].value
        if name_cell is None:
            break
        name = clean_text(name_cell)
        if name == '':
            break

        num_cell = worksheet[f'{ITEM_NUM_COL}{row}'].value
        item_num = clean_text(num_cell)

        qty_cell = worksheet[f'{ITEM_QTY_COL}{row}'].value
        quantity = clean_text(qty_cell)

        comment_cell = worksheet[f'{ITEM_COMMENT_COL}{row}'].value
        comment = clean_text(comment_cell)

        line = f"{item_num} - {name} - {quantity} шт."
        if comment:
            line += f"\n({comment})"
        items.append(line)
        logger.debug(f"Добавлена позиция: {line}")

    logger.info(f"Собрано позиций: {len(items)}")
    return items


def parse_order_excel(excel_file):
    """Сохраняет прежнюю логику разбора Excel без изменения структуры."""
    wb = openpyxl.load_workbook(excel_file, data_only=True)
    ws = wb.active

    line = clean_text(ws['B2'].value)
    client = clean_text(ws[CLIENT_CELL].value)
    manager_raw = clean_text(ws[MANAGER_CELL].value)

    if manager_raw:
        if ':' in manager_raw:
            manager = manager_raw.split(':', 1)[-1].strip()
        else:
            manager = manager_raw
    else:
        manager = ""

    logger.info(f"B2: {line}")
    logger.info(f"CLIENT ({CLIENT_CELL}): {client}")
    logger.info(f"MANAGER ({MANAGER_CELL}): {manager}")

    m = re.search(r"\D*(\d+)\D+(\d{2}\.\d{2}\.\d{4})", line)
    if not m:
        raise ValueError("ERROR parsing B2")

    return {
        'order_number': m.group(1),
        'order_date': m.group(2),
        'client': client,
        'manager': manager,
        'emails': extract_excel_emails(ws),
        'items': extract_items_fixed(ws),
    }


# ================= ФОРМИРОВАНИЕ РЕЗУЛЬТАТА =================

def build_order_title(order_number, order_date):
    if order_date:
        return f"Заказ {order_number} от {order_date}"
    return f"Заказ {order_number}"


def build_header_lines(source_type, order_title, client, manager):
    lines = [order_title, f"Заказчик: {client}"]

    # Для Excel сохраняем прежнюю строку менеджера даже при пустом значении.
    # Для TXT строка добавляется только тогда, когда менеджер реально указан.
    if source_type == 'excel' or manager:
        lines.append(f"Менеджер: {manager}")

    return lines


def sanitize_windows_name(name):
    invalid = '<>:"/\\|?*'
    for ch in invalid:
        name = name.replace(ch, "_")
    return name.strip()


def write_email_file(work_folder, found_emails):
    email_file_path = work_folder / EMAIL_FILENAME

    if found_emails:
        try:
            with open(email_file_path, 'w', encoding='utf-8') as f:
                f.write(', '.join(found_emails))
            logger.info(f"Emails saved to: {EMAIL_FILENAME}")
        except Exception as e:
            logger.error(f"Error writing email file: {e}")
    else:
        logger.info("No emails found.")
        if email_file_path.exists():
            try:
                email_file_path.unlink()
                logger.info(f"Removed old {EMAIL_FILENAME} (no emails).")
            except Exception as e:
                logger.error(f"Error removing old email file: {e}")

    return email_file_path


def write_items_file(work_folder, header_lines, items):
    items_file_path = work_folder / ITEMS_FILENAME

    if items:
        try:
            header = "\n".join(header_lines)
            content = header + "\n\n\n" + "\n\n".join(items)
            with open(items_file_path, 'w', encoding='utf-8') as f:
                f.write(content)
            logger.info(f"Items saved to: {ITEMS_FILENAME} (найдено {len(items)} позиций)")
        except Exception as e:
            logger.error(f"Error writing items file: {e}")
    else:
        logger.warning("No items found in the table.")
        if items_file_path.exists():
            try:
                items_file_path.unlink()
                logger.info(f"Removed old {ITEMS_FILENAME} (no items).")
            except Exception as e:
                logger.error(f"Error removing old items file: {e}")

    return items_file_path


def open_path(path, success_message, error_prefix):
    try:
        os.startfile(str(path))
        logger.info(success_message)
    except Exception as e:
        logger.error(f"{error_prefix}: {e}")


def main():
    folder = get_working_folder(sys.argv[1:])

    logger.info(f"Working folder: {folder}")

    if not folder.exists() or not folder.is_dir():
        fail("Target folder does not exist or is not a directory.")

    # Сохраняем прежнее распределение файлов: Excel, CDR и служебные TXT остаются
    # в корне, остальные материалы перемещаются в папку «файлы».
    files_folder = move_customer_files_to_files_folder(folder)
    logger.info(f"Files folder: {files_folder}")

    # Сначала ищем корректный TXT заказа. Это важно: в папке-шаблоне может уже
    # находиться Excel, но наличие входного TXT означает, что заказ передала
    # Наталия Миронова и должна использоваться именно TXT-ветка.
    txt_result = find_order_txt(folder)

    excel_file = None
    for f in folder.iterdir():
        if is_order_excel(f):
            excel_file = f
            break

    source_type = None
    source_txt_file = None
    source_txt_relative = None

    if txt_result:
        source_txt_file, txt_text, txt_encoding = txt_result
        source_type = 'txt'
        logger.info(f"Data source selected: TXT ({source_txt_file})")
        if excel_file:
            logger.info(
                f"Excel {excel_file.name} is present but ignored because a valid order TXT was found."
            )
        logger.info(f"TXT encoding: {txt_encoding}")

        try:
            data = parse_order_txt(txt_text)
        except Exception as e:
            fail(f"Не удалось разобрать TXT-файл {source_txt_file.name}: {e}")

        # Исходный TXT всегда должен лежать в папке «файлы», независимо от имени.
        source_txt_file = ensure_source_txt_in_files_folder(source_txt_file, files_folder)
        source_txt_relative = source_txt_file.relative_to(folder)

    elif excel_file:
        source_type = 'excel'
        logger.info(f"Data source selected: Excel ({excel_file.name})")
        try:
            data = parse_order_excel(excel_file)
        except Exception as e:
            fail(f"Ошибка разбора Excel {excel_file.name}: {e}")
    else:
        fail("Не найден ни Excel-файл заказа, ни подходящий TXT-файл с данными заказа.")

    order_number = data['order_number']
    order_date = data['order_date']
    client = data['client']
    manager = data['manager']
    found_emails = data['emails']
    items = data['items']

    # Критическая страховка: TXT-заказ всегда считается заказом Наталии Мироновой.
    # Не полагаемся только на parse_order_txt — повторно фиксируем менеджера здесь,
    # непосредственно перед созданием заголовков, имени папки и имени CDR.
    if source_type == 'txt':
        manager = TXT_MANAGER
        data['manager'] = TXT_MANAGER
        logger.info(f"TXT MODE FORCED: manager set to {TXT_MANAGER}")

    if source_type == 'txt' and manager != TXT_MANAGER:
        fail(f"Внутренняя ошибка TXT-режима: менеджер не равен {TXT_MANAGER}")

    logger.info(f"ORDER: {order_number}")
    logger.info(f"DATE: {order_date or '[нет в источнике]'}")
    logger.info(f"CLIENT: {client}")
    logger.info(f"MANAGER: {manager or '[нет в источнике]'}")
    if found_emails:
        logger.info(f"Found emails: {', '.join(found_emails)}")
    logger.info(f"Items count: {len(items)}")

    order_title = build_order_title(order_number, order_date)

    base_name = f"{order_title} ({client})"
    if source_type == 'txt':
        # В TXT-режиме имя менеджера обязательно и не зависит от содержимого файла.
        new_name_folder = f"{base_name} - {TXT_MANAGER}"
    elif manager:
        new_name_folder = f"{base_name} - {manager}"
    else:
        new_name_folder = base_name

    new_name_folder = sanitize_windows_name(new_name_folder)
    if source_type == 'txt' and not new_name_folder.endswith(f"- {TXT_MANAGER}"):
        fail(f"Внутренняя ошибка: имя TXT-заказа сформировано без менеджера: {new_name_folder}")
    logger.info(f"New folder name: {new_name_folder}")

    header_lines = build_header_lines(
        source_type=source_type,
        order_title=order_title,
        client=client,
        manager=manager,
    )

    # Еще одна явная проверка результата, чтобы TXT-заказ физически не мог
    # сформировать файл позиций без строки менеджера.
    if source_type == 'txt':
        expected_manager_line = f"Менеджер: {TXT_MANAGER}"
        header_lines = [line for line in header_lines if not line.startswith("Менеджер:")]
        header_lines.append(expected_manager_line)
        logger.info(f"TXT header forced: {expected_manager_line}")

    clipboard_text = "\r\n".join(header_lines)
    try:
        pyperclip.copy(clipboard_text)
        logger.info("Copied to clipboard:\n" + clipboard_text.replace('\r', ''))
    except Exception as e:
        logger.error(f"Clipboard error: {e}")

    # --- Переименование папки ---
    parent = folder.parent
    target_folder = parent / new_name_folder

    if target_folder == folder:
        logger.info("Folder already correctly named.")
        work_folder = folder
    elif target_folder.exists():
        target_folder = unique_path(target_folder)
        try:
            folder.rename(target_folder)
            work_folder = target_folder
            logger.info(f"Folder renamed with unique suffix: {target_folder.name}")
        except PermissionError:
            logger.error("Folder busy, cannot rename. Working inside current folder.")
            work_folder = folder
    else:
        try:
            folder.rename(target_folder)
            work_folder = target_folder
            logger.info("Folder renamed.")
        except PermissionError:
            logger.error("Folder busy, cannot rename. Working inside current folder.")
            work_folder = folder

    logger.info(f"Work folder now: {work_folder}")
    files_folder = work_folder / FILES_FOLDER_NAME
    files_folder.mkdir(exist_ok=True)
    final_order_name = work_folder.name
    if source_type == 'txt' and not final_order_name.endswith(f"- {TXT_MANAGER}"):
        fail(
            f"TXT-заказ оказался в папке без менеджера: {final_order_name}. "
            f"Ожидалось окончание: - {TXT_MANAGER}"
        )

    # После переименования папки пересобираем актуальный путь к исходному TXT,
    # а затем переименовываем его уже в окончательной папке заказа.
    source_txt_work_path = None
    if source_type == 'txt' and source_txt_relative is not None:
        source_txt_work_path = work_folder / source_txt_relative
        source_txt_work_path = rename_source_txt(source_txt_work_path)
        if source_txt_work_path.name != TXT_SOURCE_FILENAME:
            fail(
                f"Не удалось переименовать исходный TXT в {TXT_SOURCE_FILENAME}. "
                f"Текущее имя: {source_txt_work_path.name}"
            )
        logger.info(f"Final source TXT path: {source_txt_work_path}")

    # Если есть дополнительные вложения, открываем сразу папку «файлы»,
    # чтобы они не остались незамеченными. Один входной TXT вложением не считается.
    if has_additional_files(files_folder, source_txt_work_path):
        folder_to_open = files_folder
        folder_open_message = "Files folder opened in Explorer."
    else:
        folder_to_open = work_folder
        folder_open_message = "Order folder opened in Explorer."

    # --- Открытие папки в Проводнике ---
    open_path(folder_to_open, folder_open_message, "Cannot open folder")

    # --- Запись служебных файлов ---
    write_email_file(work_folder, found_emails)
    items_file_path = write_items_file(work_folder, header_lines, items)

    # Актуализация пути к Excel после возможного переименования папки.
    excel_path = work_folder / excel_file.name if excel_file else None

    # --- Переименование CDR-файла (с заменой) ---
    cdr_file = find_cdr_file(work_folder)

    if not cdr_file:
        logger.warning("No .cdr found.")

        # Сохраняем прежнее поведение Excel-ветки:
        # если CDR нет, Excel открывается здесь и еще раз в обязательном блоке ниже.
        if excel_path:
            open_path(excel_path, "Excel opened.", "Cannot open Excel")
    else:
        new_cdr = cdr_file.parent / f"{final_order_name}.cdr"
        logger.info(f"Target CDR name: {new_cdr.name}")
        if cdr_file == new_cdr:
            logger.info("CDR already has correct name.")
            target_cdr = cdr_file
        else:
            try:
                os.replace(str(cdr_file), str(new_cdr))
                logger.info("CDR renamed (replaced if existed).")
                target_cdr = new_cdr
            except PermissionError:
                logger.error("CDR file busy, opening original.")
                target_cdr = cdr_file
            except OSError as e:
                logger.error(f"Error renaming CDR: {e}")
                target_cdr = cdr_file

        open_path(target_cdr, "CorelDRAW opened.", "Cannot open CDR")

    # --- Открытие Excel (как и раньше — всегда для Excel-источника) ---
    if excel_path:
        open_path(excel_path, "Excel opened.", "Cannot open Excel")

    # --- Открытие файла с позициями заказа (если создан) ---
    if items_file_path.exists():
        open_path(
            items_file_path,
            f"Opened {ITEMS_FILENAME}",
            "Cannot open items file"
        )

    logger.info("DONE.")


if __name__ == '__main__':
    main()
