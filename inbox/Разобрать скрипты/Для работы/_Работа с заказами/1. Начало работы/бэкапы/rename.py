import logging
import os
import pathlib
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile

import openpyxl
import pyperclip

# ================= НАСТРОЙКИ =================
CLIENT_CELL = 'D4'                # ячейка с заказчиком
MANAGER_CELL = 'D5'               # ячейка с менеджером
EMAIL_FILENAME = "- Emails для согласования.txt"
ITEMS_FILENAME = "- Позиции заказа.txt"
LOG_FILENAME = "rename.log"       # файл лога (в папке скрипта)
FILES_FOLDER_NAME = "файлы"
TXT_MANAGER = "Наталия Миронова"
TXT_SOURCE_FILENAME = "заказ от Мироновой.txt"
SCRIPT_VERSION = "2026-07-24-manager-fixed-v11"

ORDERS_DIR = pathlib.Path(r"D:\YandexDisk\Проекты\Работа в Компас\Заказы")
TEMPLATE_FOLDER = ORDERS_DIR / "● [копировать] Заказ 00000 (ЗАКАЗЧИК) - МЕНЕДЖЕР"

# ---------- Настройки для позиций заказа ----------
ITEMS_START_ROW = 8               # строка первой позиции
ITEM_NUM_COL = 'C'                # колонка с номером
ITEM_NAME_COL = 'D'               # колонка с наименованием
ITEM_QTY_COL = 'E'                # колонка с количеством
ITEM_COMMENT_COL = 'F'            # колонка с комментарием
# ====================================================

EMAIL_PATTERN = r'[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}'


def clean_text(raw):
    """
    Удаляет подстроки вида _xNNNN_ (например, _x000D_),
    но оставляет реальные переносы строк.
    """
    if raw is None:
        return ""
    text = str(raw)
    text = re.sub(r'_x[0-9A-Fa-f]{4}_', '', text)
    return text.strip()


script_path = pathlib.Path(__file__).resolve()
script_dir = script_path.parent
log_file = script_dir / LOG_FILENAME

logger = logging.getLogger('rename')
logger.setLevel(logging.DEBUG)
logger.handlers.clear()

# Лог каждый запуск создается заново, без накопления старого мусора.
file_handler = logging.FileHandler(log_file, encoding='utf-8', mode='w')
file_handler.setLevel(logging.DEBUG)
formatter = logging.Formatter(
    '%(asctime)s - %(levelname)s - %(message)s',
    datefmt='%Y-%m-%d %H:%M:%S'
)
file_handler.setFormatter(formatter)
logger.addHandler(file_handler)

console_handler = logging.StreamHandler()
console_handler.setLevel(logging.INFO)
console_handler.setFormatter(formatter)
logger.addHandler(console_handler)

logger.info("=== Python folder rename started ===")
logger.info(f"SCRIPT VERSION: {SCRIPT_VERSION}")
logger.info(f"SCRIPT PATH: {script_path}")


def fail(message):
    logger.error(message)
    input("Press Enter...")
    sys.exit(1)


def normalize_arg(raw_arg):
    folder_str = raw_arg.strip().strip('"').strip()
    path = pathlib.Path(folder_str)
    if not path.is_absolute():
        path = (script_dir / path).resolve()
    return path


def unique_path(path):
    if not path.exists():
        return path

    parent = path.parent

    if path.is_dir() or not path.suffix:
        stem = path.name
        suffix = ""
    else:
        stem = path.stem
        suffix = path.suffix

    for index in range(1, 1000):
        candidate = parent / f"{stem} ({index}){suffix}"
        if not candidate.exists():
            return candidate

    raise FileExistsError(f"Cannot create unique name for {path}")


def is_archive(path):
    return path.is_file() and path.suffix.lower() in {
        ".zip", ".7z", ".rar", ".tar", ".gz", ".tgz", ".bz2", ".tbz2", ".xz", ".txz"
    }


def extract_archive(archive_path, destination):
    suffix = archive_path.suffix.lower()

    if suffix == ".zip":
        with zipfile.ZipFile(archive_path) as archive:
            archive.extractall(destination)
        return

    try:
        shutil.unpack_archive(str(archive_path), str(destination))
        return
    except (shutil.ReadError, ValueError):
        pass

    seven_zip = shutil.which("7z") or shutil.which("7za") or shutil.which("7zr")
    if seven_zip:
        subprocess.run(
            [seven_zip, "x", str(archive_path), f"-o{destination}", "-y"],
            check=True,
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
        )
        return

    raise RuntimeError(
        f"Cannot unpack {archive_path.name}. Для .rar/.7z нужен установленный 7-Zip в PATH."
    )


def move_file_or_folder(source, destination_dir):
    destination = unique_path(destination_dir / source.name)
    shutil.move(str(source), str(destination))
    logger.info(f"Moved: {source.name} -> {destination}")


def is_order_excel(path):
    return (
        path.is_file()
        and path.name.startswith('Заказ')
        and path.suffix.lower() in ['.xlsx', '.xlsm']
    )


def should_stay_in_order_root(item):
    if item.is_dir():
        return False
    if is_order_excel(item):
        return True
    if item.suffix.lower() == ".cdr":
        return True
    if item.suffix.lower() in [".txt", ".url"] and item.name.startswith("-"):
        return True
    return False


def move_customer_files_to_files_folder(order_folder):
    files_folder = order_folder / FILES_FOLDER_NAME
    files_folder.mkdir(exist_ok=True)

    for item in order_folder.iterdir():
        if item == files_folder:
            continue
        if should_stay_in_order_root(item):
            continue
        move_file_or_folder(item, files_folder)

    return files_folder


def find_cdr_file(order_folder):
    search_dirs = [order_folder, order_folder / FILES_FOLDER_NAME]
    for search_dir in search_dirs:
        if not search_dir.exists() or not search_dir.is_dir():
            continue
        for f in search_dir.iterdir():
            if f.is_file() and f.suffix.lower() == ".cdr":
                return f
    return None


def create_order_folder_from_inputs(input_paths):
    if not TEMPLATE_FOLDER.exists() or not TEMPLATE_FOLDER.is_dir():
        fail(f"Template folder does not exist: {TEMPLATE_FOLDER}")

    ORDERS_DIR.mkdir(parents=True, exist_ok=True)
    new_folder = unique_path(ORDERS_DIR / TEMPLATE_FOLDER.name)
    shutil.copytree(TEMPLATE_FOLDER, new_folder)
    logger.info(f"Template copied to: {new_folder}")

    for input_path in input_paths:
        if not input_path.exists():
            fail(f"Dragged path does not exist: {input_path}")

        if is_archive(input_path):
            logger.info(f"Extracting archive: {input_path.name}")
            extract_archive(input_path, new_folder)
            try:
                input_path.unlink()
                logger.info(f"Archive deleted: {input_path}")
            except Exception as e:
                logger.error(f"Cannot delete archive {input_path}: {e}")
        else:
            move_file_or_folder(input_path, new_folder)

    return new_folder


def get_working_folder(args):
    if not args:
        fail("No target folder, archive or files passed. Use bat file.")

    paths = [normalize_arg(arg) for arg in args]
    logger.debug(f"Raw argv: {[repr(arg) for arg in args]}")
    logger.info(f"Received paths: {len(paths)}")

    if len(paths) == 1 and paths[0].is_dir():
        return paths[0]

    return create_order_folder_from_inputs(paths)


# ================= ЧТЕНИЕ TXT =================

def read_txt_with_encoding(path):
    """
    Читает TXT в распространенных кодировках.
    Возвращает кортеж: (текст, название кодировки).
    """
    raw = path.read_bytes()

    if raw.startswith((b'\xff\xfe', b'\xfe\xff')):
        encodings = ['utf-16', 'utf-8-sig', 'utf-8', 'cp1251']
    else:
        encodings = ['utf-8-sig', 'utf-8', 'cp1251']

    last_error = None
    for encoding in encodings:
        try:
            return raw.decode(encoding), encoding
        except UnicodeDecodeError as e:
            last_error = e

    raise UnicodeError(f"Не удалось определить кодировку TXT: {path.name}. {last_error}")



def looks_like_order_txt(text):
    """Проверяет TXT. Двоеточия после заголовков необязательны."""
    has_order = re.search(
        r'(?im)^\s*заказ\s*(?:№|N)?\s*:?\s*\d+',
        text
    ) is not None
    has_client = re.search(
        r'(?im)^\s*заказчик\b\s*:?\s*.+$',
        text
    ) is not None
    has_items = re.search(
        r'(?im)^\s*состав\s+заказа\s*:?\s*$',
        text
    ) is not None
    return has_order and has_client and has_items


def find_order_txt(order_folder):
    """
    Ищет TXT с данными заказа по содержимому, а не по имени.
    Поиск идет в корне заказа, в папке «файлы» и во вложенных папках.
    Служебные выходные TXT исключаются.
    """
    ignored_names = {EMAIL_FILENAME.lower(), ITEMS_FILENAME.lower()}
    candidates = []

    for path in order_folder.rglob('*'):
        if not path.is_file() or path.suffix.lower() != '.txt':
            continue
        if path.name.lower() in ignored_names:
            continue

        try:
            text, encoding = read_txt_with_encoding(path)
        except Exception as e:
            logger.warning(f"Cannot read TXT {path}: {e}")
            continue

        if looks_like_order_txt(text):
            candidates.append((path, text, encoding))

    if not candidates:
        return None

    if len(candidates) > 1:
        names = '; '.join(str(item[0]) for item in candidates)
        fail(f"Найдено несколько TXT с данными заказа: {names}")

    return candidates[0]


def remove_inn_from_client(raw_client):
    """
    Удаляет ИНН из строки заказчика. ИНН отдельно не сохраняется.

    Поддерживаются, например:
        (ИНН: 7838490236) Сибгазполимер
        Сибгазполимер (ИНН: 7838490236)
        Сибгазполимер, ИНН 7838490236
    """
    client = clean_text(raw_client)
    client = re.sub(r'\(\s*ИНН\s*:?\s*[^)]*\)', '', client, flags=re.IGNORECASE)
    client = re.sub(r'\bИНН\s*:?\s*\d+\b', '', client, flags=re.IGNORECASE)
    client = re.sub(r'\s{2,}', ' ', client)
    return client.strip(' \t,;:—–-()')


def unique_emails(text):
    """Находит email и удаляет повторы без изменения исходного порядка."""
    result = []
    seen = set()

    for email in re.findall(EMAIL_PATTERN, text):
        normalized = email.lower()
        if normalized in seen:
            continue
        seen.add(normalized)
        result.append(email)

    return result



def parse_txt_items(text):
    """
    Извлекает позиции после «Состав заказа».
    Количество может стоять перед многострочным комментарием.
    """
    section_match = re.search(
        r'(?im)^\s*состав\s+заказа\s*:?\s*$',
        text
    )
    if not section_match:
        return []

    section = text[section_match.end():]
    raw_items = []
    current_num = None
    current_lines = []

    def flush_current():
        if current_num is not None and current_lines:
            raw_item = "\n".join(current_lines).strip()
            if raw_item:
                raw_items.append((current_num, raw_item))

    for raw_line in section.splitlines():
        line = clean_text(raw_line)
        item_match = re.match(r'^\s*(\d+)\s*[.)]\s*(.*)$', line)

        if item_match:
            flush_current()
            current_num = item_match.group(1)
            current_lines = [item_match.group(2)]
        elif current_num is not None and line:
            current_lines.append(line)

    flush_current()

    items = []
    for item_num, raw_item in raw_items:
        qty_match = re.match(
            r'(?s)^(.*?)\s*-\s*(\d+(?:[.,]\d+)?)\s*шт\.?\s*(.*)$',
            raw_item,
            flags=re.IGNORECASE
        )

        if not qty_match:
            logger.warning(
                f"Не удалось распознать количество в позиции {item_num}: {raw_item}"
            )
            continue

        name = clean_text(qty_match.group(1))
        quantity = clean_text(qty_match.group(2))
        comment_lines = [
            clean_text(part)
            for part in qty_match.group(3).splitlines()
            if clean_text(part)
        ]
        comment = "\n".join(comment_lines).strip()

        line = f"{item_num} - {name} - {quantity} шт."
        if comment:
            if comment.startswith("(") and comment.endswith(")"):
                line += f"\n{comment}"
            else:
                line += f"\n({comment})"

        items.append(line)
        logger.debug(f"Добавлена TXT-позиция: {line}")

    logger.info(f"Собрано TXT-позиций: {len(items)}")
    return items



def parse_order_txt(text):
    """Разбирает TXT. Для TXT менеджер всегда Наталия Миронова."""
    order_match = re.search(
        r'(?im)^\s*заказ\s*(?:№|N)?\s*:?\s*(\d+)'
        r'(?:\s+от\s+(\d{2}\.\d{2}\.\d{4}))?',
        text
    )
    if not order_match:
        raise ValueError("не найден номер заказа")

    order_number = order_match.group(1)
    order_date = order_match.group(2) or ""

    if not order_date:
        date_match = re.search(
            r'(?im)^\s*дата(?:\s+заказа)?\s*:?\s*'
            r'(\d{2}\.\d{2}\.\d{4})\s*$',
            text
        )
        if date_match:
            order_date = date_match.group(1)

    client_match = re.search(
        r'(?im)^\s*заказчик\b\s*:?\s*(.*?)\s*$',
        text
    )
    if not client_match:
        raise ValueError("не найден заказчик")

    client = remove_inn_from_client(client_match.group(1))
    if not client:
        raise ValueError("не удалось выделить название заказчика")

    items = parse_txt_items(text)
    if not items:
        raise ValueError("не найдены корректные позиции заказа")

    return {
        'order_number': order_number,
        'order_date': order_date,
        'client': client,
        'manager': TXT_MANAGER,
        'emails': unique_emails(text),
        'items': items,
    }


def extract_excel_emails(worksheet):
    found_emails = []

    for row in range(1, worksheet.max_row + 1):
        cell = worksheet[f'B{row}']
        if cell.value and "Контакты для согласования макетов" in str(cell.value):
            cell_text = clean_text(str(cell.value))
            found_emails = unique_emails(cell_text)
            if found_emails:
                logger.info(f"Found emails: {', '.join(found_emails)}")
            break

    return found_emails


def extract_items_fixed(worksheet):
    """
    Собирает позиции, начиная с ITEMS_START_ROW, пока в колонке ITEM_NAME_COL есть значение.
    Формат:
        основная строка: номер - наименование - количество шт.
        если есть комментарий: новая строка (в круглых скобках, с сохранением внутренних переносов).
    Между позициями позже будут добавлены пустые строки.
    """
    items = []
    for row in range(ITEMS_START_ROW, worksheet.max_row + 1):
        name_cell = worksheet[f'{ITEM_NAME_COL}{row}'].value
        if name_cell is None:
            break
        name = clean_text(name_cell)
        if name == '':
            break

        num_cell = worksheet[f'{ITEM_NUM_COL}{row}'].value
        item_num = clean_text(num_cell)

        qty_cell = worksheet[f'{ITEM_QTY_COL}{row}'].value
        quantity = clean_text(qty_cell)

        comment_cell = worksheet[f'{ITEM_COMMENT_COL}{row}'].value
        comment = clean_text(comment_cell)

        line = f"{item_num} - {name} - {quantity} шт."
        if comment:
            line += f"\n({comment})"
        items.append(line)
        logger.debug(f"Добавлена позиция: {line}")

    logger.info(f"Собрано позиций: {len(items)}")
    return items


def parse_order_excel(excel_file):
    """Сохраняет прежнюю логику разбора Excel без изменения структуры."""
    wb = openpyxl.load_workbook(excel_file, data_only=True)
    ws = wb.active

    line = clean_text(ws['B2'].value)
    client = clean_text(ws[CLIENT_CELL].value)
    manager_raw = clean_text(ws[MANAGER_CELL].value)

    if manager_raw:
        if ':' in manager_raw:
            manager = manager_raw.split(':', 1)[-1].strip()
        else:
            manager = manager_raw
    else:
        manager = ""

    logger.info(f"B2: {line}")
    logger.info(f"CLIENT ({CLIENT_CELL}): {client}")
    logger.info(f"MANAGER ({MANAGER_CELL}): {manager}")

    m = re.search(r"\D*(\d+)\D+(\d{2}\.\d{2}\.\d{4})", line)
    if not m:
        raise ValueError("ERROR parsing B2")

    return {
        'order_number': m.group(1),
        'order_date': m.group(2),
        'client': client,
        'manager': manager,
        'emails': extract_excel_emails(ws),
        'items': extract_items_fixed(ws),
    }


# ================= ФОРМИРОВАНИЕ РЕЗУЛЬТАТА =================

def build_order_title(order_number, order_date):
    if order_date:
        return f"Заказ {order_number} от {order_date}"
    return f"Заказ {order_number}"


def build_header_lines(source_type, order_title, client, manager):
    lines = [order_title, f"Заказчик: {client}"]

    # Для Excel сохраняем прежнюю строку менеджера даже при пустом значении.
    # Для TXT строка добавляется только тогда, когда менеджер реально указан.
    if source_type == 'excel' or manager:
        lines.append(f"Менеджер: {manager}")

    return lines


def sanitize_windows_name(name):
    invalid = '<>:"/\\|?*'
    for ch in invalid:
        name = name.replace(ch, "_")
    return name.strip()


def write_email_file(work_folder, found_emails):
    email_file_path = work_folder / EMAIL_FILENAME

    if found_emails:
        try:
            with open(email_file_path, 'w', encoding='utf-8') as f:
                f.write(', '.join(found_emails))
            logger.info(f"Emails saved to: {EMAIL_FILENAME}")
        except Exception as e:
            logger.error(f"Error writing email file: {e}")
    else:
        logger.info("No emails found.")
        if email_file_path.exists():
            try:
                email_file_path.unlink()
                logger.info(f"Removed old {EMAIL_FILENAME} (no emails).")
            except Exception as e:
                logger.error(f"Error removing old email file: {e}")

    return email_file_path


def write_items_file(work_folder, header_lines, items):
    items_file_path = work_folder / ITEMS_FILENAME

    if items:
        try:
            header = "\n".join(header_lines)
            content = header + "\n\n\n" + "\n\n".join(items)
            with open(items_file_path, 'w', encoding='utf-8') as f:
                f.write(content)
            logger.info(f"Items saved to: {ITEMS_FILENAME} (найдено {len(items)} позиций)")
        except Exception as e:
            logger.error(f"Error writing items file: {e}")
    else:
        logger.warning("No items found in the table.")
        if items_file_path.exists():
            try:
                items_file_path.unlink()
                logger.info(f"Removed old {ITEMS_FILENAME} (no items).")
            except Exception as e:
                logger.error(f"Error removing old items file: {e}")

    return items_file_path


def open_path(path, success_message, error_prefix):
    try:
        os.startfile(str(path))
        logger.info(success_message)
    except Exception as e:
        logger.error(f"{error_prefix}: {e}")


def select_and_parse_source(root_folder):
    """
    Ищет источник и полностью проверяет его.
    TXT имеет приоритет, потому что TXT-заказы передаёт Наталия Миронова.
    """
    txt_result = find_order_txt(root_folder)
    if txt_result:
        txt_path, txt_text, txt_encoding = txt_result
        data = parse_order_txt(txt_text)
        data['manager'] = TXT_MANAGER
        logger.info(
            f"Preflight source: TXT ({txt_path.name}), encoding: {txt_encoding}"
        )
        return 'txt', data

    excel_files = sorted(
        [path for path in root_folder.rglob('*') if is_order_excel(path)],
        key=lambda path: str(path).lower()
    )
    if excel_files:
        data = parse_order_excel(excel_files[0])
        logger.info(f"Preflight source: Excel ({excel_files[0].name})")
        return 'excel', data

    raise ValueError(
        "не найден ни корректный Excel-файл заказа, "
        "ни подходящий TXT-файл с данными заказа"
    )


def preflight_inputs(input_paths):
    """
    Проверяет заказ ДО копирования шаблона и перемещения исходников.
    При ошибке папка заказа не создаётся.
    """
    for path in input_paths:
        if not path.exists():
            raise FileNotFoundError(f"Dragged path does not exist: {path}")

    if len(input_paths) == 1 and input_paths[0].is_dir():
        return select_and_parse_source(input_paths[0])

    with tempfile.TemporaryDirectory(prefix="rename_preflight_") as temp_dir:
        scan_root = pathlib.Path(temp_dir)

        for index, input_path in enumerate(input_paths, start=1):
            if input_path.is_dir():
                shutil.copytree(
                    input_path,
                    scan_root / f"folder_{index}"
                )
            elif is_archive(input_path):
                archive_folder = scan_root / f"archive_{index}"
                archive_folder.mkdir()
                extract_archive(input_path, archive_folder)
            elif input_path.suffix.lower() in {'.txt', '.xlsx', '.xlsm'}:
                shutil.copy2(
                    input_path,
                    unique_path(scan_root / input_path.name)
                )

        return select_and_parse_source(scan_root)


def rename_source_txt(source_txt_file):
    """Переименовывает исходный TXT, не затирая существующий файл."""
    target = source_txt_file.with_name(TXT_SOURCE_FILENAME)

    if source_txt_file == target:
        return source_txt_file

    if target.exists():
        previous = unique_path(
            target.with_name(f"{target.stem} - предыдущий{target.suffix}")
        )
        target.rename(previous)

    source_txt_file.rename(target)
    logger.info(
        f"Source TXT renamed: {source_txt_file.name} -> {TXT_SOURCE_FILENAME}"
    )
    return target


def has_additional_files(files_folder, source_txt_file=None):
    """
    Для TXT сам исходный TXT не считается дополнительным вложением.
    Любой другой файл или папка означает, что нужно открыть папку «файлы».
    """
    source_resolved = (
        source_txt_file.resolve()
        if source_txt_file is not None
        else None
    )

    for path in files_folder.iterdir():
        if source_resolved is not None and path.resolve() == source_resolved:
            continue
        return True

    return False


def main():
    if len(sys.argv) < 2:
        fail("No target folder, archive or files passed. Use bat file.")

    input_paths = [normalize_arg(arg) for arg in sys.argv[1:]]
    logger.debug(f"Raw argv: {[repr(arg) for arg in sys.argv[1:]]}")
    logger.info(f"Received paths: {len(input_paths)}")

    # Сначала полностью проверяем источник. До успешной проверки
    # шаблон не копируется и исходные файлы не перемещаются.
    try:
        source_type, data = preflight_inputs(input_paths)
    except Exception as e:
        fail(f"Предварительная проверка не пройдена: {e}")

    folder = get_working_folder(sys.argv[1:])
    logger.info(f"Working folder: {folder}")

    if not folder.exists() or not folder.is_dir():
        fail("Target folder does not exist or is not a directory.")

    files_folder = move_customer_files_to_files_folder(folder)
    logger.info(f"Files folder: {files_folder}")

    order_number = data['order_number']
    order_date = data['order_date']
    client = data['client']
    found_emails = data['emails']
    items = data['items']

    # КРИТИЧНО: для TXT менеджер задаётся здесь повторно,
    # непосредственно перед формированием всех имён и заголовков.
    if source_type == 'txt':
        manager = TXT_MANAGER
        data['manager'] = TXT_MANAGER
    else:
        manager = data['manager']

    if source_type == 'txt' and manager != TXT_MANAGER:
        fail("Внутренняя ошибка: у TXT-заказа потерялся менеджер")

    logger.info(f"ORDER: {order_number}")
    logger.info(f"DATE: {order_date or '[нет в источнике]'}")
    logger.info(f"CLIENT: {client}")
    logger.info(f"MANAGER: {manager}")

    order_title = build_order_title(order_number, order_date)
    base_name = f"{order_title} ({client})"

    if source_type == 'txt':
        new_name_folder = f"{base_name} - {TXT_MANAGER}"
    elif manager:
        new_name_folder = f"{base_name} - {manager}"
    else:
        new_name_folder = base_name

    new_name_folder = sanitize_windows_name(new_name_folder)

    if source_type == 'txt' and not new_name_folder.endswith(
        f"- {TXT_MANAGER}"
    ):
        fail("Внутренняя ошибка: имя папки сформировано без менеджера")

    header_lines = [
        order_title,
        f"Заказчик: {client}",
        f"Менеджер: {TXT_MANAGER if source_type == 'txt' else manager}"
    ]

    if source_type == 'txt' and header_lines[-1] != (
        f"Менеджер: {TXT_MANAGER}"
    ):
        fail("Внутренняя ошибка: заголовок позиций без менеджера")

    logger.info(f"New folder name: {new_name_folder}")

    try:
        pyperclip.copy("\r\n".join(header_lines))
    except Exception as e:
        logger.error(f"Clipboard error: {e}")

    parent = folder.parent
    target_folder = parent / new_name_folder

    if target_folder == folder:
        work_folder = folder
    elif target_folder.exists():
        target_folder = unique_path(target_folder)
        try:
            folder.rename(target_folder)
            work_folder = target_folder
        except PermissionError:
            logger.error(
                "Folder busy, cannot rename. Working inside current folder."
            )
            work_folder = folder
    else:
        try:
            folder.rename(target_folder)
            work_folder = target_folder
        except PermissionError:
            logger.error(
                "Folder busy, cannot rename. Working inside current folder."
            )
            work_folder = folder

    logger.info(f"Work folder now: {work_folder}")
    final_order_name = work_folder.name
    files_folder = work_folder / FILES_FOLDER_NAME
    files_folder.mkdir(exist_ok=True)

    if source_type == 'txt' and not final_order_name.endswith(
        f"- {TXT_MANAGER}"
    ):
        fail(
            f"TXT-заказ оказался в папке без менеджера: {final_order_name}"
        )

    source_txt_file = None
    if source_type == 'txt':
        txt_result = find_order_txt(work_folder)
        if not txt_result:
            fail("После перемещения не найден исходный TXT заказа")
        source_txt_file = rename_source_txt(txt_result[0])

    write_email_file(work_folder, found_emails)
    items_file_path = write_items_file(
        work_folder,
        header_lines,
        items
    )

    cdr_file = find_cdr_file(work_folder)
    if cdr_file:
        new_cdr = cdr_file.parent / f"{final_order_name}.cdr"

        if source_type == 'txt' and not new_cdr.stem.endswith(
            f"- {TXT_MANAGER}"
        ):
            fail("Внутренняя ошибка: имя CDR сформировано без менеджера")

        if cdr_file != new_cdr:
            try:
                os.replace(str(cdr_file), str(new_cdr))
            except PermissionError:
                logger.error("CDR file busy, opening original.")
                new_cdr = cdr_file
            except OSError as e:
                logger.error(f"Error renaming CDR: {e}")
                new_cdr = cdr_file

        open_path(new_cdr, "CorelDRAW opened.", "Cannot open CDR")
    else:
        logger.warning("No .cdr found.")

    if source_type == 'excel':
        excel_files = [
            path for path in work_folder.iterdir()
            if is_order_excel(path)
        ]
        if excel_files:
            open_path(
                excel_files[0],
                "Excel opened.",
                "Cannot open Excel"
            )

    if has_additional_files(files_folder, source_txt_file):
        open_path(
            files_folder,
            "Files folder opened in Explorer.",
            "Cannot open files folder"
        )
    else:
        open_path(
            work_folder,
            "Order folder opened in Explorer.",
            "Cannot open order folder"
        )

    if items_file_path.exists():
        open_path(
            items_file_path,
            f"Opened {ITEMS_FILENAME}",
            "Cannot open items file"
        )

    logger.info("DONE.")


if __name__ == '__main__':
    main()
